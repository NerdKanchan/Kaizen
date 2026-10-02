"""Excel / CSV exchange for relationships so quality teams can maintain terminology in the tool they already use."""

import csv
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from kaizen.models import Relationship
from kaizen.reporting.workbook import csv_text, save_workbook
from kaizen.terminology.repository import TerminologyRepository

COLUMNS = ["ID", "Canonical", "Aliases", "Scope", "Doc Types", "Item Anchors", "Provenance", "Created By", "Active", "Notes", "Version", "Created At", "Updated At"]
_SEP = "; "


@dataclass
class ImportResult:
    created: int = 0
    updated: int = 0
    unchanged: int = 0
    errors: list[str] = field(default_factory=list)

    def summary(self) -> str:
        return f"created {self.created}, updated {self.updated}, unchanged {self.unchanged}, errors {len(self.errors)}"


def _rel_to_row(r: Relationship) -> list[Any]:
    return [r.id, r.canonical, _SEP.join(r.aliases), r.scope, _SEP.join(d.value for d in r.doc_types), _SEP.join(r.item_anchors), r.provenance, r.created_by, "Y" if r.active else "N", r.notes, r.version, r.created_at.isoformat(timespec="seconds"), r.updated_at.isoformat(timespec="seconds")]


def _split(v: Any) -> list[str]:
    if v is None:
        return []
    return [x.strip() for x in str(v).replace(";", "\n").replace(",", "\n").split("\n") if x.strip()] if "\n" in str(v) or ";" in str(v) else [x.strip() for x in str(v).split(",") if x.strip()]


def _split_aliases(v: Any) -> list[str]:
    # aliases may legitimately contain commas ("TAPE, SURGICAL"); only ';' or newlines separate entries
    if v is None:
        return []
    return [x.strip() for x in str(v).replace("\n", ";").split(";") if x.strip()]


def export_xlsx(repo: TerminologyRepository, path: Path | str) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Relationships"
    ws.append(COLUMNS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3864")
    for r in repo.list(active_only=False):
        ws.append(_rel_to_row(r))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:M{max(ws.max_row, 2)}"
    for col, width in zip("ABCDEFGHIJKLM", (10, 40, 60, 18, 16, 16, 12, 14, 8, 50, 9, 22, 22)):
        ws.column_dimensions[col].width = width
    info = wb.create_sheet("How to edit")
    for line in (
        "Edit rows on the Relationships sheet and import the file with `kaizen terminology import <file>`.",
        "Leave ID blank to create a new relationship; keep the ID to update an existing one (a new version is recorded).",
        "Aliases and Item Anchors are separated by semicolons. Scope is 'global', 'family:<prefix>' or 'sku:<code>'.",
        "Doc Types (optional): BOM, LABEL, DRAWING, PCO separated by semicolons; blank means all document types.",
        "Active: Y or N. Version, Created At and Updated At are informational and ignored on import.",
    ):
        info.append([line])
    info.column_dimensions["A"].width = 120
    save_workbook(wb, path)
    return path


def export_csv(repo: TerminologyRepository, path: Path | str) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(COLUMNS)
        for r in repo.list(active_only=False):
            w.writerow([csv_text(value) for value in _rel_to_row(r)])
    return path


def import_xlsx(repo: TerminologyRepository, path: Path | str, imported_by: str = "import") -> ImportResult:
    """Import the native column format through the same validation used by the CLI and UI."""
    from kaizen.terminology.source_import import import_source

    result = import_source(repo, path, imported_by)
    return ImportResult(result.created if result.applied else 0, result.updated if result.applied else 0, result.unchanged, result.errors)


def import_csv(repo: TerminologyRepository, path: Path | str, imported_by: str = "import") -> ImportResult:
    return import_xlsx(repo, path, imported_by)
