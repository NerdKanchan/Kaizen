"""Preview and import readable relationship exports with explicit column mapping."""

import csv
import hashlib
import re
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from kaizen.ingest.workbook import workbook_access_problem
from kaizen.matching.normalize import normalize
from kaizen.models import Relationship
from kaizen.terminology.exchange import COLUMNS, _split, _split_aliases
from kaizen.terminology.repository import TerminologyRepository
from kaizen.terminology.store import RelationshipStore

FIELDS = COLUMNS[:10]
MULTI_FIELDS = {"Aliases", "Item Anchors", "Doc Types"}


@dataclass
class SourceTable:
    sheets: list[str]
    sheet: str
    header_row: int
    columns: list[str]
    rows: list[tuple[int, dict[str, str]]]
    sha256: str
    header_error: str | None = None


def _cell_text(cell) -> str:
    value = cell.value
    if value is None:
        return ""
    # Excel often stores item numbers numerically, with the leading zeros in the display format.
    if isinstance(value, (int, float)) and not isinstance(value, bool) and float(value).is_integer():
        text = str(int(value))
        if re.fullmatch(r"0+", cell.number_format or ""):
            return text.zfill(len(cell.number_format))
        return text
    return str(value).strip()


def read_source(path: Path | str, sheet: str | None = None, header_row: int = 1, *, strict_header: bool = True) -> SourceTable:
    path = Path(path)
    if not 1 <= header_row <= 1000:
        raise ValueError("Header row must be between 1 and 1000.")
    if path.suffix.lower() not in (".xlsx", ".xlsm", ".csv"):
        raise ValueError("Supply a readable XLSX, XLSM or CSV relationship export.")
    if problem := workbook_access_problem(path):
        raise ValueError(problem)
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    if path.suffix.lower() == ".csv":
        if sheet not in (None, "", "CSV"):
            raise ValueError("CSV files have one table named CSV.")
        with path.open(newline="", encoding="utf-8-sig") as fh:
            raw = list(csv.reader(fh))
        sheets, selected = ["CSV"], "CSV"
    else:
        try:
            wb = load_workbook(path, read_only=True, data_only=False)
        except Exception as e:
            raise ValueError("Could not read the workbook. Supply a readable XLSX or CSV export.") from e
        try:
            sheets = wb.sheetnames
            selected = sheet or ("Relationships" if "Relationships" in sheets else sheets[0])
            if selected not in sheets:
                raise ValueError(f"Worksheet {selected!r} does not exist.")
            raw = [[_cell_text(c) for c in row] for row in wb[selected].iter_rows()]
        finally:
            wb.close()
    columns = [str(v).strip() for v in raw[header_row - 1]] if len(raw) >= header_row else []
    named = [c for c in columns if c]
    error = "The selected header row does not exist." if len(raw) < header_row else "The selected header row has no column names." if not named else None
    if len(named) != len(set(named)):
        error = "Column names must be unique. Select the correct header row or rename duplicate headers in the export."
    if error:
        if strict_header:
            raise ValueError(error)
        return SourceTable(sheets, selected, header_row, [], [], digest, error)
    rows = []
    for number, values in enumerate(raw[header_row:], start=header_row + 1):
        row = {name: str(values[i]).strip() if i < len(values) else "" for i, name in enumerate(columns) if name}
        if any(row.values()):
            rows.append((number, row))
    return SourceTable(sheets, selected, header_row, named, rows, digest)


def inspect_source(path: Path | str, sheet: str | None = None, header_row: int = 1) -> dict[str, Any]:
    table = read_source(path, sheet, header_row, strict_header=False)
    # Only suggest Kaizen's known field names. Source terminology meanings require an explicit mapping.
    known = {f.casefold(): f for f in FIELDS}
    mapping = {known[c.casefold()]: [c] for c in table.columns if c.casefold() in known}
    return {"sheets": table.sheets, "sheet": table.sheet, "header_row": table.header_row, "columns": table.columns,
            "row_count": len(table.rows), "sample": [{"row": n, "values": r} for n, r in table.rows[:5]], "column_map": mapping,
            "sha256": table.sha256, "header_error": table.header_error}


@dataclass
class SourceImportResult:
    created: int = 0
    updated: int = 0
    unchanged: int = 0
    errors: list[str] = field(default_factory=list)
    dry_run: bool = False
    applied: bool = False
    rows: list[dict[str, Any]] = field(default_factory=list)
    source: dict[str, Any] = field(default_factory=dict)

    def summary(self) -> str:
        prefix = "Preview" if self.dry_run else "Import blocked; no changes saved" if self.errors else "Import"
        return f"{prefix}: created {self.created}, updated {self.updated}, unchanged {self.unchanged}, errors {len(self.errors)}"

    def to_dict(self) -> dict[str, Any]:
        return asdict(self) | {"summary": self.summary()}


def _mapping(columns: list[str], column_map: dict[str, Any] | None) -> dict[str, list[str]]:
    if column_map is None:
        known = {f.casefold(): f for f in FIELDS}
        column_map = {known[c.casefold()]: [c] for c in columns if c.casefold() in known}
    if not isinstance(column_map, dict):
        raise ValueError("Column mapping must be an object mapping Kaizen fields to source column names.")
    mapping = {}
    for name, selected in column_map.items():
        if name not in FIELDS:
            raise ValueError(f"Unknown relationship field {name!r}.")
        selected = [selected] if isinstance(selected, str) else selected
        if not isinstance(selected, list) or any(not isinstance(c, str) for c in selected):
            raise ValueError(f"Mapping for {name} must contain column names.")
        if not selected:
            continue
        if name not in MULTI_FIELDS and len(selected) != 1:
            raise ValueError(f"Select one source column for {name}.")
        for c in selected:
            if c not in columns:
                raise ValueError(f"Source column {c!r} does not exist.")
        mapping[name] = list(dict.fromkeys(selected))
    if not mapping.get("Canonical"):
        raise ValueError("Select the canonical wording column, or use a Kaizen terminology export with a Canonical column.")
    return mapping


def _signature(rel: Relationship) -> tuple:
    return (normalize(rel.canonical).sorted_key, tuple(sorted({normalize(a).sorted_key for a in rel.aliases})),
            rel.scope, tuple(sorted(d.value for d in rel.doc_types)), tuple(sorted(rel.item_anchors)))


def import_source(
    repo: TerminologyRepository, path: Path | str, imported_by: str = "import", *,
    column_map: dict[str, Any] | None = None, sheet: str | None = None, header_row: int = 1,
    default_scope: str = "global", dry_run: bool = False, source_name: str | None = None, expected_version: str | None = None,
) -> SourceImportResult:
    """Plan all rows before saving anything. Counts describe planned changes when previewed or blocked.

    Existing IDs are updated only when mapped explicitly. Rows without IDs are deduplicated by their
    terms, scope, document types and anchors; importing another file never reactivates an existing rule.
    """
    table = read_source(path, sheet, header_row)
    mapping = _mapping(table.columns, column_map)
    default_scope = Relationship(id="validate", canonical="validate", scope=default_scope).scope
    result = SourceImportResult(dry_run=dry_run, source={"file": source_name or Path(path).name, "sha256": table.sha256,
                                                      "sheet": table.sheet, "header_row": header_row, "column_map": mapping})
    native = mapping["Canonical"][0].casefold() == "canonical" and all(c.casefold() == "aliases" for c in mapping.get("Aliases", []))
    if not native and not (mapping.get("Aliases") or mapping.get("Item Anchors") or mapping.get("ID")):
        raise ValueError("Map equivalent wording or an item anchor alongside the canonical wording.")
    with repo.db.transaction():
        existing = repo.list(active_only=False)
        version = RelationshipStore(existing).snapshot().version
        result.source["terminology_version"] = version
        if expected_version is not None and expected_version != version:
            raise ValueError("Terminology changed since the preview. Preview the import again before saving.")
        by_id = {r.id: r for r in existing}
        original_ids = set(by_id)
        by_signature = {_signature(r): r for r in existing}
        next_number = int(repo.next_id().split("-")[1])
        reserved = set(by_id)
        source_ids: set[str] = set()
        planned: list[tuple[str, Relationship, dict[str, Any], int]] = []
        for number, row in table.rows:
            mapped = {name: "; ".join(row[c] for c in selected if row[c]) for name, selected in mapping.items()}
            if not any(mapped.values()):
                continue
            try:
                if any(row[c].startswith("=") for selected in mapping.values() for c in selected):
                    raise ValueError("Mapped cells contain formulas. Export their evaluated values before importing.")
                if not mapped.get("Canonical", "").strip():
                    raise ValueError("Canonical wording is missing.")
                active = mapped.get("Active", "").upper()
                if active not in ("", "Y", "YES", "TRUE", "1", "ACTIVE", "N", "NO", "FALSE", "0", "INACTIVE"):
                    raise ValueError("Active must be Y or N (or true/false).")
                rel_id = mapped.get("ID", "").strip()
                if rel_id and rel_id in source_ids:
                    raise ValueError(f"ID {rel_id} appears more than once in the source.")
                if rel_id:
                    source_ids.add(rel_id)
                current = by_id.get(rel_id)
                if rel_id and rel_id not in original_ids and rel_id in reserved:
                    raise ValueError(f"ID {rel_id} conflicts with another row in the source.")
                if rel_id and current is None and repo.history(rel_id):
                    raise ValueError(f"ID {rel_id} has deleted history; use a new ID.")
                aliases = list(dict.fromkeys(_split_aliases(mapped.get("Aliases"))))
                canonical = mapped["Canonical"]
                if not native:
                    aliases = [a for a in aliases if a.casefold() != canonical.casefold()]
                fields = {"canonical": canonical, "aliases": aliases, "scope": mapped.get("Scope") or default_scope,
                          "doc_types": list(dict.fromkeys(d.upper() for d in _split(mapped.get("Doc Types")))),
                          "item_anchors": list(dict.fromkeys(_split(mapped.get("Item Anchors")))), "notes": mapped.get("Notes", ""),
                          "active": active not in ("N", "NO", "FALSE", "0", "INACTIVE")}
                if current:
                    for name, attr in {"Aliases": "aliases", "Scope": "scope", "Doc Types": "doc_types", "Item Anchors": "item_anchors", "Notes": "notes", "Active": "active"}.items():
                        if name not in mapping:
                            fields[attr] = getattr(current, attr)
                rel = Relationship(id=rel_id or "pending", **fields)
                if not native and not (rel.aliases or rel.item_anchors):
                    raise ValueError("No distinct equivalent wording or item anchor was supplied.")
                duplicate = by_signature.get(_signature(rel)) if not rel_id else None
                if duplicate:
                    action, rel = "unchanged", duplicate
                elif current:
                    changed = {k: v for k, v in fields.items() if getattr(current, k) != getattr(rel, k)}
                    action = "update" if changed else "unchanged"
                    rel = rel.model_copy(update={"version": current.version + (action == "update")})
                else:
                    if not rel_id:
                        while f"REL-{next_number:03d}" in reserved:
                            next_number += 1
                        rel_id = f"REL-{next_number:03d}"
                        next_number += 1
                    reserved.add(rel_id)
                    origin = f"Source: {result.source['file']}; SHA-256 {table.sha256}; sheet {table.sheet}; row {number}."
                    fields["notes"] = "\n".join(x for x in (fields["notes"], origin) if x)
                    rel = Relationship(id=rel_id, provenance=mapped.get("Provenance") or "imported", created_by=imported_by, **fields)
                    action, changed = "create", fields
                if action in ("create", "update"):
                    if current:
                        old_signature = _signature(current)
                        indexed = by_signature.get(old_signature)
                        if indexed and indexed.id == current.id:
                            by_signature.pop(old_signature)
                            replacement = next((r for r in by_id.values() if r.id != current.id and _signature(r) == old_signature), None)
                            if replacement:
                                by_signature[old_signature] = replacement
                    by_id[rel.id] = rel
                    by_signature[_signature(rel)] = rel
                planned.append((action, rel, changed if action in ("create", "update") else {}, number))
                result.rows.append({"row": number, "action": action, "id": rel.id, "canonical": rel.canonical, "aliases": rel.aliases,
                                    "scope": rel.scope, "item_anchors": rel.item_anchors, "doc_types": [d.value for d in rel.doc_types], "active": rel.active})
                counter = {"create": "created", "update": "updated", "unchanged": "unchanged"}[action]
                setattr(result, counter, getattr(result, counter) + 1)
            except ValueError as e:
                result.errors.append(f"{table.sheet} row {number}: {e}")
        if not table.rows or not result.rows and not result.errors:
            result.errors.append("No relationship rows found in the selected table and mapping.")
        if dry_run or result.errors:
            return result
        for action, rel, fields, number in planned:
            if action == "create":
                repo.create(rel_id=rel.id, provenance=rel.provenance, created_by=imported_by, **fields)
            elif action == "update":
                origin = f"imported from {result.source['file']}; SHA-256 {table.sha256}; sheet {table.sheet}; row {number}"
                repo.update(rel.id, changed_by=imported_by, change_note=origin, **fields)
        result.applied = True
    return result
