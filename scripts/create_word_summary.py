from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("output/docx/kaizen-project-summary.docx")


FONT = "Liberation Sans"


def set_font(run, name=FONT, size=None, bold=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), fill)
    tc_pr.append(shade)


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_body(doc, text, after=3):
    p = doc.add_paragraph(style="Body")
    p.paragraph_format.space_after = Pt(after)
    p.add_run(text)
    return p


def add_heading(doc, text):
    p = doc.add_paragraph(style="Section")
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(2)
    p.add_run(text)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="Bullet")
    p.paragraph_format.space_after = Pt(1)
    p.add_run("- " + text)
    return p


def main():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.58)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(10.25)
    normal.font.color.rgb = RGBColor(25, 35, 48)

    body = doc.styles.add_style("Body", WD_STYLE_TYPE.PARAGRAPH)
    body.base_style = normal
    body.font.name = FONT
    body._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    body._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    body.font.size = Pt(10.25)
    body.paragraph_format.line_spacing = 1.05

    section_style = doc.styles.add_style("Section", WD_STYLE_TYPE.PARAGRAPH)
    section_style.base_style = normal
    section_style.font.name = FONT
    section_style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    section_style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    section_style.font.size = Pt(11.5)
    section_style.font.bold = True
    section_style.font.color.rgb = RGBColor(0, 0, 0)

    bullet = doc.styles.add_style("Bullet", WD_STYLE_TYPE.PARAGRAPH)
    bullet.base_style = body
    bullet.paragraph_format.left_indent = Inches(0.20)
    bullet.paragraph_format.first_line_indent = Inches(-0.15)
    bullet.paragraph_format.line_spacing = 1.0

    title_style = doc.styles["Title"]
    title_style.font.name = FONT
    title_style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    title_style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    title_style_ppr = title_style.element.get_or_add_pPr()
    title_style_border = title_style_ppr.find(qn("w:pBdr"))
    if title_style_border is not None:
        title_style_ppr.remove(title_style_border)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_after = Pt(2)
    r = title.add_run("Kaizen BOM Cross Check Automation Project Summary")
    set_font(r, name=FONT, size=17, bold=True, color=(0, 0, 0))
    p_pr = title._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is not None:
        p_pr.remove(p_bdr)

    sub = doc.add_paragraph(style="Body")
    sub.paragraph_format.space_after = Pt(7)
    r = sub.add_run("Innovation Hackathon 2026 | Project response to the BOM Cross Check Automation challenge")
    set_font(r, size=9.5, color=(80, 80, 80))

    add_heading(doc, "What problem we are solving")
    add_body(doc, "Sustenance Engineering projects require teams to confirm that the Bill of Materials (BOM), product label, packaging drawing and Product Change Order (PCO) all describe the same kit. Today, the BOM is downloaded from JDE while labels, drawings and PCOs come from MasterControl. A project reviewer and an independent reviewer compare the files manually, mark the source documents and maintain a separate tracker. The work is repetitive, difficult to scale and vulnerable to missed discrepancies or rework.")
    add_body(doc, "The supplied challenge estimates about 60 minutes of review per SKU for each reviewer. A 129-SKU project therefore needs about 258 reviewer hours, while the independent reviewer may only be available for four hours a day. The challenge asks for a way to reduce this effort without changing JDE, MasterControl or the existing quality approval workflow.")

    add_heading(doc, "What we built")
    add_body(doc, "Kaizen Cross Check is a local review application that reads the documents reviewers already download: JDE BOMs in PDF, Excel or CSV; product labels; packaging drawings; and PCO forms. It extracts component descriptions, item numbers, quantities and source evidence, then performs the comparisons that the problem statement requires.")

    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Inches(2.05)
    table.columns[1].width = Inches(4.85)
    hdr = table.rows[0].cells
    for cell, text in zip(hdr, ("Required comparison", "Kaizen output")):
        set_cell_shading(cell, "1F4E79")
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        rr = p.add_run(text)
        set_font(rr, size=9.25, bold=True, color=(255, 255, 255))
    rows = [
        ("BOM to Label", "Confirms component presence, quantities and the label REF against the BOM parent."),
        ("BOM to Drawing and Label to Drawing", "Checks that physical components are represented consistently in the packaging drawing and label."),
        ("PCO to BOM and old to new label", "Checks approved additions, deletions, substitutions and unexpected label changes; also flags PCO affected codes without a BOM."),
    ]
    for idx, (left, right) in enumerate(rows):
        cells = table.add_row().cells
        for cell, text in zip(cells, (left, right)):
            set_cell_margins(cell)
            if idx % 2 == 1:
                set_cell_shading(cell, "F2F6FA")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            rr = p.add_run(text)
            set_font(rr, size=8.75, bold=(cell == cells[0]))

    add_heading(doc, "How Kaizen handles the hard part")
    add_body(doc, "The same component often appears under different names across documents. Kaizen first normalises descriptions and checks exact matches. It then uses editable, versioned terminology relationships so that a reviewer can define an approved equivalence for a product family or a specific item. Similarity and optional AI suggestions can identify a possible match, but they cannot approve one. Those rows remain visible for reviewer validation.")
    add_body(doc, "Every result is classified as exact, equivalent, potential, mismatch or missing. Each finding carries its source file, file hash, page, bounding box and extracted text so the reviewer can see where the result came from. The application supports independent two-reviewer decisions, blind review for the second reviewer, disagreement handling, action items and a corrected rerun to verify resolution.")

    add_heading(doc, "What the reviewer receives and why it matters")
    add_bullet(doc, "A review queue that separates clear matches from rows that need human judgement.")
    add_bullet(doc, "An audit-grade Excel workbook with compared values, classifications, discrepancies, terminology used and reviewer decisions.")
    add_bullet(doc, "An annotated BOM PDF and action-item list for follow-up, with a run comparison to show what a corrected document set resolved.")
    add_body(doc, "On the project’s seeded golden dataset, Kaizen evaluates its results against declared ground truth and reports 100 percent precision and recall. The current business-case view shows a 27 percent reduction on the first run and a 73 percent projection after reviewers approve recurring terminology relationships. The latter is a projection, not yet a measured production result. The next step is to validate extraction and effort savings using real redacted document sets.", after=0)

    core = doc.core_properties
    core.title = "Kaizen BOM Cross Check Automation Project Summary"
    core.subject = "One page project summary for the Innovation Hackathon"
    core.author = "Project Kaizen"
    doc.save(OUT)


if __name__ == "__main__":
    main()
