from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


DOCX = Path("output/docx/kaizen-project-summary.docx")
FONT = "Liberation Sans"


def set_font(run, size=None, bold=None, color=None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    node = OxmlElement("w:shd")
    node.set(qn("w:fill"), fill)
    tc_pr.append(node)


def cell_margins(cell, top=100, start=110, bottom=100, end=110):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        item = margins.find(qn(f"w:{edge}"))
        if item is None:
            item = OxmlElement(f"w:{edge}")
            margins.append(item)
        item.set(qn("w:w"), str(value))
        item.set(qn("w:type"), "dxa")


def border(cell, color="D9D9D9"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        elem = borders.find(qn(f"w:{edge}"))
        if elem is None:
            elem = OxmlElement(f"w:{edge}")
            borders.append(elem)
        elem.set(qn("w:val"), "single")
        elem.set(qn("w:sz"), "4")
        elem.set(qn("w:color"), color)


def add_section_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    set_font(r, size=11.5, bold=True, color=(0, 0, 0))
    return p


def add_body(doc, text, after=4):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.08
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text)
    set_font(r, size=10.25, color=(25, 35, 48))
    return p


def add_table(doc, rows):
    table = doc.add_table(rows=1, cols=3)
    table.autofit = False
    table.columns[0].width = Inches(3.35)
    table.columns[1].width = Inches(1.20)
    table.columns[2].width = Inches(2.35)
    headers = ("Budget item", "Amount", "Use")
    for cell, text in zip(table.rows[0].cells, headers):
        shade(cell, "1F4E79")
        cell_margins(cell)
        border(cell, "D9D9D9")
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text)
        set_font(r, size=9.25, bold=True, color=(255, 255, 255))
    for index, values in enumerate(rows):
        cells = table.add_row().cells
        for cell, text in zip(cells, values):
            cell_margins(cell)
            border(cell)
            if index % 2:
                shade(cell, "F2F6FA")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(text)
            set_font(r, size=9.25, bold=(values[0] == "Total"))
            if cell == cells[1]:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    return table


def main():
    doc = Document(DOCX)
    doc.add_page_break()

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(2)
    r = title.add_run("Budget Request for Drawing OCR Improvement")
    set_font(r, size=16, bold=True, color=(0, 0, 0))

    sub = doc.add_paragraph()
    sub.paragraph_format.space_after = Pt(8)
    r = sub.add_run("Kaizen BOM Cross Check Automation")
    set_font(r, size=9.5, color=(80, 80, 80))

    add_body(doc, "We request budget to improve extraction from packaging drawings. The current manual OCR approach is not accurate enough for dense drawings with callouts, technical labels and mixed-language text. Better extraction is needed before drawing-to-BOM and drawing-to-label checks can be relied on at scale.")

    add_section_heading(doc, "Minimum proof of concept request")
    add_table(doc, [
        ("AI model credits for drawing OCR training and evaluation", "₹3,000", "Test vision and OCR models on representative drawings."),
        ("Domain adaptation and deployment", "₹3,000", "Prepare drawing-specific prompts, rules and a deployable review flow."),
        ("Infrastructure", "₹1,500", "Storage, compute and run logging for the pilot."),
        ("Total", "₹7,500", "Minimum requested allocation."),
    ])

    add_section_heading(doc, "Recommendation")
    add_body(doc, "₹7,500 is sufficient only for a small proof of concept. It can show whether a vision-based approach improves the current OCR output, but it leaves little room for repeated testing, annotated sample preparation, accuracy measurement and deployment hardening. We recommend approving a ₹25,000 pilot budget if the objective is a reliable decision on whether the solution should progress beyond the hackathon.")

    add_section_heading(doc, "Recommended pilot budget")
    add_table(doc, [
        ("AI model credits and accuracy evaluation", "₹12,000", "Run OCR and vision-model experiments across representative drawing pages."),
        ("Domain-specific training and validation", "₹7,000", "Create evaluation samples and tune extraction for callouts and technical terminology."),
        ("Deployment and monitoring", "₹3,500", "Integrate the selected approach into Kaizen and measure extraction outcomes."),
        ("Infrastructure", "₹2,500", "Compute, storage, logging and backup during the pilot."),
        ("Total", "₹25,000", "Recommended pilot allocation."),
    ])

    add_section_heading(doc, "Expected outcome")
    add_body(doc, "The pilot will measure whether drawing extraction is accurate enough to support Kaizen's existing comparison and reviewer workflow. Funding will be used to compare model output against reviewed drawing samples, identify which callout types remain unreliable and decide whether to adopt the approach for a broader deployment. This request does not assume a production rollout; the production budget should be estimated only after pilot accuracy, document volume and operating requirements are known.", after=0)
    doc.save(DOCX)


if __name__ == "__main__":
    main()
