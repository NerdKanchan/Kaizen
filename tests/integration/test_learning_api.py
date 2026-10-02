"""Closed-testing authorization, automatic evaluation and independent verification over HTTP."""

import openpyxl
import pytest
from fastapi.testclient import TestClient

from kaizen.api import app as app_module
from kaizen.api.app import create_app
from kaizen.evaluation.learning import shadow_evaluate
from kaizen.review.access import AccessStore
from kaizen.review.auth import LocalAccounts
from kaizen.review.learning import DrawingInput, LabelApproval
from tests.unit.test_learning import AUTHOR, VERIFIER, body_for, row_for, verify
from tests.unit.test_learning import learning_env as learning_fixture

learning_env = learning_fixture

PW = 'test-password-2026'
VIEWER = 'viewer@bd.com'
OUTSIDER = 'outsider@bd.com'


@pytest.fixture
def clients(learning_env):
    ws, run, learning = learning_env
    access = AccessStore(ws.db)
    access.bootstrap(VERIFIER, PW)
    accounts = LocalAccounts(ws.db)
    app = create_app(ws, accounts=accounts)
    out = {}
    for email in (AUTHOR, VERIFIER, VIEWER, OUTSIDER):
        client = TestClient(app)
        if email != VERIFIER:
            client.post('/api/auth/signup', json={'email': email, 'password': PW})
            access.update(email, VERIFIER, 'approved', False)
        assert client.post('/api/auth/signin', json={'email': email, 'password': PW}).status_code == 200
        out[email] = client
    access.share(run.metadata.run_id, VERIFIER, 'edit', AUTHOR)
    access.share(run.metadata.run_id, VIEWER, 'view', AUTHOR)
    return out


def test_learning_authorization(learning_env, clients):
    _, run, learning = learning_env
    row = row_for(run, learning)
    base = f'/api/runs/{run.metadata.run_id}/learning'
    assert clients[OUTSIDER].get(base).status_code == 404
    assert clients[VIEWER].get(base).status_code == 200
    assert clients[VIEWER].put(f'{base}/rows/{row.row_id}', json=body_for(row).model_dump()).status_code == 403
    assert clients[AUTHOR].post(f'{base}/rows/{row.row_id}/approve', json={'expected_version': 1, 'approve': True}).status_code == 403
    assert clients[AUTHOR].post(f'{base}/shadow').status_code == 403
    assert clients[AUTHOR].get(f'{base}/export.json').status_code == 403
    assert clients[VERIFIER].post(f'{base}/enable').status_code == 403


def test_verified_feedback_automatically_builds_candidate(learning_env, clients):
    ws, run, learning = learning_env
    row = row_for(run, learning)
    base = f'/api/runs/{run.metadata.run_id}/learning'
    body = body_for(row).model_dump()
    assert clients[AUTHOR].put(f'{base}/rows/{row.row_id}', json=body).status_code == 200
    assert clients[VERIFIER].get(f'{base}/export.json').json()['samples'] == []
    response = clients[VERIFIER].post(f'{base}/rows/{row.row_id}/approve', json={'expected_version': 1, 'approve': True})
    assert response.status_code == 200, response.text
    summary = clients[AUTHOR].get(base).json()
    assert summary['automatic_learning']['state'] == 'complete'
    assert summary['latest_evaluation']['training_samples'] == 1
    assert summary['latest_evaluation']['baseline']['pairing_accuracy'] is None
    assert summary['counts']['verified_train'] == 1
    exported = clients[VERIFIER].get(f'{base}/export.json').json()
    assert len(exported['samples']) == 1
    assert clients[VERIFIER].get(f'{base}/export.json?split=evaluation').json()['samples'] == []
    # Ordinary work remains captured, but changing the evidence invalidates the label.
    response = clients[AUTHOR].post(f'/api/runs/{run.metadata.run_id}/decisions', json={'row_id': row.row_id, 'decision': 'NEEDS_MORE_INFORMATION', 'expected_revision': 0})
    assert response.status_code == 200
    assert clients[VERIFIER].get(f'{base}/export.json').json()['samples'] == []
    assert clients[VERIFIER].get(f'{base}/observations.json').json()['rows'][0]['history'][-1]['reviewer'] == AUTHOR
    assert ws.db.conn.execute('SELECT COUNT(*) FROM learning_evaluations').fetchone()[0] == 2


def test_excel_review_import_rebuilds_candidate_and_dry_run_does_not(learning_env, clients, tmp_path):
    ws, run, learning = learning_env
    row = row_for(run, learning)
    verify(learning, run, row)
    shadow_evaluate(ws, run, VERIFIER)
    base = f'/api/runs/{run.metadata.run_id}'
    path = tmp_path / 'review.xlsx'
    response = clients[AUTHOR].get(f'{base}/export.xlsx')
    assert response.status_code == 200
    path.write_bytes(response.content)
    wb = openpyxl.load_workbook(path)
    sheet = wb['BOM_Label']
    cols = {c.value: c.column for c in sheet[1]}
    target = next(n for n in range(2, sheet.max_row + 1) if sheet.cell(n, cols['Row ID']).value == row.row_id)
    sheet.cell(target, cols['Reviewer Decision'], 'NEEDS_MORE_INFORMATION')
    wb.save(path)
    report_id = learning.summary(run)['latest_evaluation']['id']
    def upload(dry):
        return clients[AUTHOR].post(f'{base}/decisions/import', data={'dry_run': str(dry).lower()}, files={'file': ('review.xlsx', path.read_bytes(), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
    assert upload(True).status_code == 200
    assert learning.summary(run)['latest_evaluation']['id'] == report_id
    assert len(learning.dataset(run, 'train')['samples']) == 1
    response = upload(False)
    assert response.status_code == 200, response.text
    assert row.row_id in response.json()['applied']
    summary = clients[AUTHOR].get(f'{base}/learning').json()
    assert summary['automatic_learning']['state'] == 'complete'
    assert summary['latest_evaluation']['training_samples'] == 0
    assert summary['latest_evaluation']['id'] != report_id


def test_bulk_review_changes_rebuild_candidate(learning_env, clients):
    ws, run, learning = learning_env
    row = row_for(run, learning, check='BOM_DRAWING')
    drawing = next(d for d in run.documents if d.id == row.source_b.doc_id)
    learning.drawing(run, row.sku, DrawingInput(expected_version=0, doc_id=drawing.id, applicability='CONFIRMED_RELEASED', released_revision='A', release_reference='Released source record', note='Verified against released drawing.'), AUTHOR)
    learning.approve('learning_drawings', run, row.sku, LabelApproval(expected_version=1, approve=True), VERIFIER)
    verify(learning, run, row, body_for(row, verdict='CORRECT_PAIR', classification='EXACT', discrepancies=[]))
    shadow_evaluate(ws, run, VERIFIER)
    assert learning.summary(run)['latest_evaluation']['training_samples'] == 1
    response = clients[AUTHOR].post(f'/api/runs/{run.metadata.run_id}/bulk-accept', json={})
    assert response.status_code == 200 and response.json()['accepted'] > 0
    assert clients[AUTHOR].get(f'/api/runs/{run.metadata.run_id}/learning').json()['automatic_learning']['state'] == 'complete'
    assert learning.summary(run)['latest_evaluation']['training_samples'] == 0


def test_reserved_sku_cannot_create_terminology(learning_env, clients):
    _, run, learning = learning_env
    row = row_for(run, learning, split='evaluation')
    response = clients[VERIFIER].post(f'/api/runs/{run.metadata.run_id}/relationships/from-row', json={'row_id': row.row_id})
    assert response.status_code == 400
    assert 'reserved for evaluation' in response.json()['detail']


def test_row_approval_preserves_verified_pairing_evidence(learning_env, clients):
    ws, run, learning = learning_env
    row = row_for(run, learning)
    base = f'/api/runs/{run.metadata.run_id}'
    assert clients[AUTHOR].post(f'{base}/decisions', json={'row_id': row.row_id, 'decision': 'CONFIRM_DISCREPANCY', 'expected_revision': 0}).status_code == 200
    verify(learning, run, row, body_for(row, review_revision=learning.review.revision(run.metadata.run_id, row.row_id)))
    shadow_evaluate(ws, run, VERIFIER)
    assert learning.summary(run)['latest_evaluation']['training_samples'] == 1
    report_id = learning.summary(run)['latest_evaluation']['id']
    response = clients[VERIFIER].post(f'{base}/finalize', json={'row_id': row.row_id, 'final_decision': 'CONFIRM_DISCREPANCY', 'expected_revision': learning.review.revision(run.metadata.run_id, row.row_id)})
    assert response.status_code == 200, response.text
    assert learning.summary(run)['latest_evaluation']['training_samples'] == 1
    assert learning.summary(run)['latest_evaluation']['id'] == report_id
    assert learning.summary(run)['latest_evaluation']['stale'] is False


def test_optimistic_conflict_and_invalid_payloads(learning_env, clients):
    _, run, learning = learning_env
    row = row_for(run, learning)
    url = f'/api/runs/{run.metadata.run_id}/learning/rows/{row.row_id}'
    body = body_for(row).model_dump()
    assert clients[AUTHOR].put(url, json=body).status_code == 200
    assert clients[AUTHOR].put(url, json=body).status_code == 409
    assert clients[AUTHOR].put(url, json={**body, 'annotated_by': VERIFIER}).status_code == 422


def test_pending_rebuild_resumes_after_application_restart(learning_env):
    ws, run, learning = learning_env
    verify(learning, run, row_for(run, learning))
    ws.db.conn.execute("INSERT INTO learning_jobs (run_id,generation,state,requested_by,updated_at) VALUES (?,1,'running',?,'2026-10-01')", (run.metadata.run_id, VERIFIER))
    ws.db.conn.commit()
    with TestClient(create_app(ws, accounts=LocalAccounts(ws.db))):
        pass
    job = ws.db.conn.execute('SELECT * FROM learning_jobs WHERE run_id=?', (run.metadata.run_id,)).fetchone()
    assert job['state'] == 'complete'
    assert job['completed_generation'] == 1
    assert learning.summary(run)['latest_evaluation']['training_samples'] == 1


def test_background_failure_is_visible_and_retryable(learning_env, clients, monkeypatch):
    _, run, learning = learning_env
    row = row_for(run, learning)
    base = f'/api/runs/{run.metadata.run_id}/learning'
    def fail(*args):
        raise RuntimeError('Simulated worker failure')
    monkeypatch.setattr(app_module, 'shadow_evaluate', fail)
    assert clients[AUTHOR].put(f'{base}/rows/{row.row_id}', json=body_for(row).model_dump()).status_code == 200
    assert clients[VERIFIER].post(f'{base}/rows/{row.row_id}/approve', json={'expected_version': 1, 'approve': True}).status_code == 200
    status = clients[AUTHOR].get(base).json()['automatic_learning']
    assert status['state'] == 'failed' and status['error']
    monkeypatch.setattr(app_module, 'shadow_evaluate', shadow_evaluate)
    retry = clients[VERIFIER].post(f'{base}/shadow')
    assert retry.status_code == 200, retry.text
    assert retry.json()['training_samples'] == 1
