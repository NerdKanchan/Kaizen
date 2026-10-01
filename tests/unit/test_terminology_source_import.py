"""Source exports must be reviewable, repeatable and traceable before they affect matching."""

import pytest
from openpyxl import Workbook

from kaizen.models import DocType
from kaizen.terminology.repository import TerminologyRepository
from kaizen.terminology.source_import import import_source, inspect_source


@pytest.fixture
def repo(tmp_path):
    return TerminologyRepository(tmp_path / "kaizen.db")


@pytest.fixture
def source(tmp_path):
    wb = Workbook()
    wb.active.title = "Read me"
    wb.active.append(["Exported terminology"])
    ws = wb.create_sheet("Approved pairs")
    ws.append(["Relationship source"])
    ws.append(["ERP Description", "Label Wording", "Drawing Wording", "Component", "Applies to"])
    ws.append(["CAP DEAD END", "End Cap", "Dead end cap", 1234, "sku:1175108DNS"])
    ws["D3"].number_format = "0000000"
    ws.append(["TAPE, SURGICAL", "Surgical Tape", "Tape", "PK001", "global"])
    path = tmp_path / "authorized.xlsx"
    wb.save(path)
    return path


MAPPING = {"Canonical": "Label Wording", "Aliases": ["ERP Description", "Drawing Wording"], "Item Anchors": "Component", "Scope": "Applies to"}


def test_inspection_and_preview_do_not_change_terminology(repo, source):
    inspection = inspect_source(source, "Approved pairs", 2)
    assert inspection["sheets"] == ["Read me", "Approved pairs"]
    assert inspection["sample"][0]["row"] == 3
    assert inspection["sample"][0]["values"]["Component"] == "0001234"
    assert inspection["column_map"] == {}, "Unknown source semantics must not be guessed."
    preview = import_source(repo, source, column_map=MAPPING, sheet="Approved pairs", header_row=2, dry_run=True)
    assert preview.created == 2 and not preview.errors and not preview.applied
    assert preview.rows[1]["aliases"] == ["TAPE, SURGICAL", "Tape"]
    assert repo.list() == [] and repo.history("REL-001") == []


def test_apply_records_source_and_uses_scoped_rule_then_reimport_is_unchanged(repo, source):
    preview = import_source(repo, source, column_map=MAPPING, sheet="Approved pairs", header_row=2, dry_run=True)
    result = import_source(repo, source, "quality@bd.com", column_map=MAPPING, sheet="Approved pairs", header_row=2,
                           expected_version=preview.source["terminology_version"])
    assert result.created == 2 and result.applied
    rel = repo.get("REL-001")
    assert rel.item_anchors == ["0001234"] and rel.created_by == "quality@bd.com" and rel.provenance == "imported"
    assert source.name in rel.notes and result.source["sha256"] in rel.notes and "row 3" in rel.notes
    hit = repo.store().lookup("CAP DEAD END", "End Cap", sku="1175108DNS", doc_types=(DocType.BOM, DocType.LABEL))
    assert hit and hit.relationship.id == rel.id
    assert repo.store().lookup("CAP DEAD END", "End Cap", sku="1275108DNS") is None
    again = import_source(repo, source, column_map=MAPPING, sheet="Approved pairs", header_row=2)
    assert again.created == again.updated == 0 and again.unchanged == 2
    assert len(repo.list()) == 2 and len(repo.history("REL-001")) == 1


def test_invalid_row_blocks_entire_import(repo, tmp_path):
    path = tmp_path / "invalid.csv"
    path.write_text("Canonical,Aliases,Scope,Active\nCap,END CAP,global,Y\nGown,GOWN ISOLATION,planet:mars,Y\n,Mask,global,Y\n")
    result = import_source(repo, path)
    assert len(result.errors) == 2 and not result.applied
    assert "row 3" in result.errors[0] and "row 4" in result.errors[1]
    assert repo.list() == []


def test_formulas_and_duplicate_headers_are_reported(repo, tmp_path):
    path = tmp_path / "formula.csv"
    path.write_text('Canonical,Aliases\nCap,"=A2"\n')
    result = import_source(repo, path)
    assert result.errors and "formulas" in result.errors[0] and repo.list() == []
    path.write_text("Canonical,Canonical\nCap,End Cap\n")
    assert "unique" in inspect_source(path)["header_error"]
    with pytest.raises(ValueError, match="unique"):
        import_source(repo, path)


def test_protected_workbook_has_actionable_error(repo, tmp_path):
    path = tmp_path / "protected.xlsx"
    path.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + "DRMEncryptedTransform".encode("utf-16le"))
    with pytest.raises(ValueError, match="authorized user.*Excel"):
        import_source(repo, path)
    assert repo.list() == []


def test_preview_refuses_apply_if_terminology_changed(repo, source):
    preview = import_source(repo, source, column_map=MAPPING, sheet="Approved pairs", header_row=2, dry_run=True)
    repo.create(canonical="Changed during review", created_by="another reviewer")
    with pytest.raises(ValueError, match="changed since the preview"):
        import_source(repo, source, column_map=MAPPING, sheet="Approved pairs", header_row=2,
                      expected_version=preview.source["terminology_version"])
    assert len(repo.list()) == 1


def test_duplicate_id_and_reserved_ids_cannot_overwrite_another_row(repo, tmp_path):
    path = tmp_path / "conflict.csv"
    path.write_text("ID,Canonical,Aliases\n,Cap,END CAP\nREL-001,Tape,SURGICAL TAPE\n")
    result = import_source(repo, path)
    assert result.errors and "conflicts" in result.errors[0] and not result.applied
    assert repo.list() == []


def test_native_updates_preserve_version_history_and_dedup_does_not_reactivate(repo, tmp_path):
    current = repo.create(canonical="Cap", aliases=["END CAP"], active=False, notes="reviewed")
    path = tmp_path / "update.csv"
    path.write_text("Canonical,Aliases\nCap,END CAP\n")
    result = import_source(repo, path)
    assert result.unchanged == 1 and not repo.get(current.id).active
    path.write_text(f"ID,Canonical,Aliases,Notes,Active\n{current.id},Cap,END CAP;DEAD END,approved change,Y\n")
    result = import_source(repo, path, "quality")
    assert result.updated == 1 and result.applied
    updated = repo.get(current.id)
    assert updated.version == 2 and updated.active
    assert repo.get_version(current.id, 1).aliases == ["END CAP"]
    assert result.source["sha256"] in repo.history(current.id)[-1].change_note


def test_unmapped_fields_are_preserved_on_id_update(repo, tmp_path):
    current = repo.create(canonical="Cap", aliases=["END CAP"], scope="sku:1175108DNS", doc_types=["BOM", "LABEL"], item_anchors=["0000123"], active=False, notes="reviewed")
    path = tmp_path / "narrow.csv"
    path.write_text(f"ID,Canonical\n{current.id},Closure\n")
    result = import_source(repo, path)
    assert result.applied and result.updated == 1
    updated = repo.get(current.id)
    assert updated.canonical == "Closure" and updated.aliases == current.aliases and updated.scope == current.scope
    assert updated.item_anchors == current.item_anchors and updated.doc_types == current.doc_types
    assert not updated.active and updated.notes == "reviewed"


def test_empty_first_sheet_still_exposes_other_worksheets(tmp_path):
    wb = Workbook()
    wb.active.title = "Cover"
    wb.create_sheet("Relationships").append(["Canonical", "Aliases"])
    path = tmp_path / "cover.xlsx"
    wb.save(path)
    inspection = inspect_source(path, sheet="Cover")
    assert inspection["sheets"] == ["Cover", "Relationships"] and inspection["header_error"]


def test_batch_deduplication_uses_the_planned_updates(repo, tmp_path):
    current = repo.create(canonical="Cap", aliases=["END CAP"])
    path = tmp_path / "batch.csv"
    path.write_text(f"ID,Canonical,Aliases\n{current.id},Closure,OTHER CLOSURE\n,Cap,END CAP\n,Closure,OTHER CLOSURE\n")
    result = import_source(repo, path)
    assert result.applied and result.updated == result.created == result.unchanged == 1
    assert {r.canonical for r in repo.list()} == {"Cap", "Closure"}
    assert result.rows[-1]["id"] == current.id
