"""Regressions found during the repository-wide closed-testing review."""

import csv
import sys
import time
from concurrent.futures import ThreadPoolExecutor
from threading import Barrier
from types import SimpleNamespace

import openpyxl
import pytest
from pydantic import ValidationError

from kaizen.api.payloads import Credentials
from kaizen.matching.semantic import sentence_transformer_embedder
from kaizen.models import MatchLevel
from kaizen.reporting.excel import export_with_review
from kaizen.reporting.workbook import save_workbook
from kaizen.review import sessions
from kaizen.review.collaborative import CollaborativeReviewStore
from kaizen.terminology.exchange import export_csv, export_xlsx
from kaizen.workspace import Workspace
from tests.unit.test_learning import AUTHOR, row_for
from tests.unit.test_learning import learning_env as learning_fixture

learning_env = learning_fixture


def test_concurrent_wrong_passwords_cannot_lose_lockout_attempts(tmp_path, monkeypatch):
    ws = Workspace(tmp_path / "ws")
    users = sessions.UserStore(ws.db)
    users.create(AUTHOR, "test-password")
    barrier = Barrier(sessions.MAX_FAILED + 2)

    def wrong(*args):
        time.sleep(0.005)  # expose a read/check/write race
        return False

    monkeypatch.setattr(sessions, "verify_password", wrong)

    def attempt(_):
        barrier.wait()
        return users.verify(AUTHOR, "wrong-password")

    try:
        with ThreadPoolExecutor(max_workers=barrier.parties) as pool:
            assert not any(pool.map(attempt, range(barrier.parties)))
        assert users._row(AUTHOR)["failed_attempts"] == sessions.MAX_FAILED
        assert users.locked_seconds(AUTHOR) > 0
    finally:
        ws.db.close()


def test_duplicate_concurrent_signup_is_a_normal_conflict(tmp_path):
    ws = Workspace(tmp_path / "ws")
    users = sessions.UserStore(ws.db)
    barrier = Barrier(5)

    def create(_):
        barrier.wait()
        try:
            users.create(AUTHOR, "test-password")
            return "created"
        except KeyError:
            return "conflict"

    try:
        with ThreadPoolExecutor(max_workers=barrier.parties) as pool:
            results = list(pool.map(create, range(barrier.parties)))
        assert results.count("created") == 1
        assert results.count("conflict") == 4
    finally:
        ws.db.close()


@pytest.mark.parametrize("changes", [{"email": "x" * 321}, {"password": "x" * 1025}])
def test_public_credentials_have_input_bounds(changes):
    with pytest.raises(ValidationError):
        Credentials.model_validate({"email": AUTHOR, "password": "test-password", **changes})


def test_exports_preserve_formula_like_comments_as_literal_text(learning_env, tmp_path):
    ws, run, learning = learning_env
    row = row_for(run, learning)
    comment = '=HYPERLINK("https://example.invalid", "source")'
    CollaborativeReviewStore(ws.db).decide(run.metadata.run_id, row.row_id, 1, AUTHOR, "ACCEPT", comment)
    wb = openpyxl.load_workbook(export_with_review(ws, run, tmp_path / "report.xlsx"))
    comments = [cell for sheet in wb for cells in sheet for cell in cells if cell.value == comment]
    assert len(comments) == 2  # current decision and immutable approval history
    assert all(cell.data_type == "s" for cell in comments)
    assert all(cell.data_type != "f" for sheet in wb for cells in sheet for cell in cells)
    assert wb["BOM_Label"].data_validations.dataValidation
    assert wb["BOM_Label"].conditional_formatting


def test_worklist_export_excludes_reserved_evaluation_rows(learning_env, tmp_path):
    ws, run, learning = learning_env
    for row in run.results:
        if row.check.value == "BOM_LABEL" and row.source_a and row.source_b and row.source_a.description == "NEEDLE 22 GA":
            row.match_level = MatchLevel.FUZZY
            row.source_a.description = "NEEDLE 22 GA LONG"
    wb = openpyxl.load_workbook(export_with_review(ws, run, tmp_path / "report.xlsx"))
    sheet = wb["Terminology_Worklist"]
    headers = {cell.value: cell.column for cell in sheet[4]}
    row_ids = {rid.strip() for cells in sheet.iter_rows(min_row=5) for rid in str(cells[headers["Row IDs"] - 1].value or "").split(",") if rid.strip()}
    assert row_ids
    reserved = {r.row_id for r in run.results if learning.split(r.sku) == "evaluation"}
    assert not row_ids.intersection(reserved)


def test_relationship_exports_store_business_text_safely(tmp_path):
    ws = Workspace(tmp_path / "ws")
    try:
        ws.repository.create(canonical="+NEEDLE", aliases=["@NEEDLE"], notes="=1+1")
        wb = openpyxl.load_workbook(export_xlsx(ws.repository, tmp_path / "terms.xlsx"))
        row = next(cells for cells in wb["Relationships"].iter_rows(min_row=2) if cells[1].value == "+NEEDLE")
        assert row[9].value == "=1+1" and row[9].data_type == "s"
        with export_csv(ws.repository, tmp_path / "terms.csv").open(newline="") as file:
            csv_row = next(r for r in csv.DictReader(file) if r["Canonical"] == "'+NEEDLE")
        assert csv_row["Aliases"] == "'@NEEDLE" and csv_row["Notes"] == "'=1+1"
    finally:
        ws.db.close()


def test_failed_export_preserves_previous_file_and_removes_tempfile(tmp_path):
    target = tmp_path / "report.xlsx"
    target.write_bytes(b"previous export")

    class BrokenWorkbook:
        def __iter__(self):
            return iter(())

        def save(self, path):
            path.write_bytes(b"incomplete new export")
            raise OSError("disk failure")

    with pytest.raises(OSError, match="disk failure"):
        save_workbook(BrokenWorkbook(), target)
    assert target.read_bytes() == b"previous export"
    assert list(tmp_path.iterdir()) == [target]


def test_semantic_embedder_requires_local_model_files(monkeypatch):
    calls = []

    class LocalModel:
        def __init__(self, name, **kwargs):
            calls.append((name, kwargs))

        def encode(self, texts, **kwargs):
            return SimpleNamespace(tolist=lambda: [[1.0, 0.0] for _ in texts])

    monkeypatch.setitem(sys.modules, "sentence_transformers", SimpleNamespace(SentenceTransformer=LocalModel))
    assert sentence_transformer_embedder("/models/approved")(["needle"]) == [[1.0, 0.0]]
    assert calls == [("/models/approved", {"local_files_only": True})]
